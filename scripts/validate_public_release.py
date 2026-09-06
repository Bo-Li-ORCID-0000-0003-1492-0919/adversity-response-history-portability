#!/usr/bin/env python3
from __future__ import annotations

import csv
import json
import re
import sys
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader
from decimal import Decimal

REPO = Path(__file__).resolve().parents[1]
FORBIDDEN_EXTENSIONS = {".dta", ".sav", ".sas7bdat", ".parquet", ".feather", ".pkl", ".pickle", ".rds", ".rdata", ".mat", ".h5", ".hdf5", ".npy", ".npz", ".db", ".sqlite"}
ID_HEADERS = {"pidp", "hhid", "hhidpn", "person_id", "participant_id", "episode_id", "target_episode_id"}
ERRORS = ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")
EXPECTED_ROWS = {
    "Figure_1_source.csv": 10, "Figure_2_source.csv": 12, "Figure_3_source.csv": 24, "Figure_4_source.csv": 36,
    **{f"Extended_Data_Figure_{i}_source.csv": n for i, n in enumerate([40, 28, 4, 56, 52, 56, 64, 48], 1)},
}
SAMPLE_PERSON_COUNT_COLUMNS = {
    "persons", "target_rows", "history_rows", "preregistration_expected_persons",
    "current_pre1_complete", "pre2_complete", "age_known", "sex_known",
    "education_known", "participants_with_shared_history_observations",
    "n_parent", "n_complete", "missing_excluded", "n_evaluated",
    "n_finite_model_predictions", "temporal_train_n", "temporal_test_n",
    "temporal_person_overlap_excluded",
}
RESULT_AUXILIARY_PERSON_COUNT_COLUMNS = {
    "n_parent", "n_complete", "missing_excluded", "n_evaluated",
    "n_finite_model_predictions", "temporal_train_n", "temporal_test_n",
    "temporal_person_overlap_excluded",
}
DIRECT_COUNT_COLUMNS = SAMPLE_PERSON_COUNT_COLUMNS | {"n_persons", "n_episodes"}
AI_TOOL_PATTERN = re.compile(
    r"\b(?:" + "chat" + "gpt|" + "co" + "dex|" + "open" + "ai)",
    re.IGNORECASE,
)
LOCAL_PATH_PATTERN = re.compile(
    "/" + "(?:Users|Volumes)/|" + "Synology" + "Drive|C:" + r"\\Users\\",
    re.IGNORECASE,
)
SECRET_PATTERNS = (
    re.compile(r"-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----"),
    re.compile(r"\bghp_[A-Za-z0-9]{30,}\b"),
    re.compile(r"\bgithub_pat_[A-Za-z0-9_]{30,}\b"),
    re.compile(r"\bAKIA[0-9A-Z]{16}\b"),
    re.compile(r"\bsk-[A-Za-z0-9]{20,}\b"),
)


def fail(message: str) -> None:
    raise AssertionError(message)


def check_files() -> None:
    ignored_parts = {".git", "__MACOSX", "__pycache__", ".ipynb_checkpoints", "outputs"}
    ignored_names = {".DS_Store"}
    files = []
    for path in REPO.rglob("*"):
        if not path.is_file():
            continue
        if any(part in ignored_parts for part in path.relative_to(REPO).parts):
            continue
        if path.name in ignored_names or path.name.startswith("._"):
            continue
        files.append(path)
    for path in files:
        rel = path.relative_to(REPO)
        if path.stat().st_size == 0:
            fail(f"empty file: {rel}")
        if path.stat().st_size > 100 * 1024 * 1024:
            fail(f"file exceeds 100 MB: {rel}")
        if path.suffix.lower() in FORBIDDEN_EXTENSIONS:
            fail(f"restricted data extension: {rel}")
        if path.name in {"paths.local.yaml", "README_BLOCKERS.txt"}:
            fail(f"excluded file: {rel}")
        if path.suffix.lower() in {".py", ".md", ".txt", ".yaml", ".yml", ".json", ".csv", ".cff", ".command", ".sh"}:
            text = path.read_text(encoding="utf-8", errors="replace")
            if LOCAL_PATH_PATTERN.search(text):
                fail(f"non-generic local path: {rel}")
            if any(pattern.search(text) for pattern in SECRET_PATTERNS):
                fail(f"possible credential or private key: {rel}")


def check_csv() -> None:
    for path in sorted((REPO / "source_data").rglob("*.csv")):
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
        headers = {str(item).strip().lower() for item in (rows[0].keys() if rows else [])}
        found = sorted(headers & ID_HEADERS)
        if found:
            fail(f"identifier fields in {path.relative_to(REPO)}: {found}")
        if path.name in EXPECTED_ROWS and len(rows) != EXPECTED_ROWS[path.name]:
            fail(f"row count mismatch in {path.name}: {len(rows)}")
        if path.name in EXPECTED_ROWS:
            ids = [int(float(row["display_row"])) for row in rows]
            if len(ids) != len(set(ids)):
                fail(f"duplicate display_row in {path.name}")


def _records(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def _integer_count(raw: object) -> int | None:
    text = str(raw or "").strip()
    if not text:
        return None
    try:
        value = Decimal(text)
    except Exception:
        return None
    if value != value.to_integral_value():
        return None
    return int(value)


def check_display_provenance() -> None:
    display_files = sorted((REPO / "source_data/main").glob("*.csv")) + sorted((REPO / "source_data/extended_data").glob("*.csv"))
    for path in display_files:
        for record in _records(path):
            source = REPO / record["source_file"]
            if not source.is_file():
                fail(f"missing provenance source: {record['source_file']}")
            original = _records(source)[int(record["source_row_0based"])]
            for key, value in json.loads(record["source_filter"]).items():
                if str(original[key]) != str(value):
                    fail(f"provenance filter mismatch: {path.relative_to(REPO)} row {record['display_row']} {key}")
            for target, column_key in (("estimate", "source_estimate_column"), ("n", "source_n_column")):
                value = record.get(target, "")
                column = record.get(column_key, "")
                if not value or not column:
                    continue
                tolerance = Decimal("0") if target == "n" else Decimal("1e-14")
                if abs(Decimal(value) - Decimal(original[column])) > tolerance:
                    fail(f"provenance value mismatch: {path.relative_to(REPO)} row {record['display_row']} {target}")
            ci_columns = [item for item in record.get("source_ci_columns", "").split(";") if item]
            for target, column in zip(("ci_lower", "ci_upper"), ci_columns):
                value = record.get(target, "")
                if value and abs(Decimal(value) - Decimal(original[column])) > Decimal("1e-14"):
                    fail(f"provenance interval mismatch: {path.relative_to(REPO)} row {record['display_row']} {target}")


def check_small_counts() -> None:
    for path in sorted((REPO / "source_data/analysis_outputs").rglob("*.csv")):
        rows = _records(path)
        if not rows:
            continue
        for row_number, row in enumerate(rows, start=2):
            for column in DIRECT_COUNT_COLUMNS & set(row):
                value = _integer_count(row.get(column, ""))
                if value is not None and 1 <= value <= 10:
                    fail(f"unsuppressed small count in {path.relative_to(REPO)}:{row_number} {column}={value}")

            if str(row.get("count_suppressed", "")).strip().lower() == "true":
                if not str(row.get("public_release_note", "")).strip():
                    fail(f"suppressed count lacks a release note: {path.relative_to(REPO)}:{row_number}")
                if path.name in {"04_recency_benchmark_results.csv", "05_sensitivity_results.csv"}:
                    exposed = sorted(
                        column for column in RESULT_AUXILIARY_PERSON_COUNT_COLUMNS
                        if str(row.get(column, "")).strip()
                    )
                    if exposed:
                        fail(
                            f"suppressed sample retains auxiliary participant counts in "
                            f"{path.relative_to(REPO)}:{row_number}: {exposed}"
                        )

    episode_path = REPO / "source_data/analysis_outputs/NMH_confirmatory/02_event_episode_counts.csv"
    rows = _records(episode_path)
    for cohort in {row["cohort"] for row in rows}:
        episode_rows = [
            row for row in rows
            if row["cohort"] == cohort
            and row["record_type"] == "episode_status"
            and row["category"] in {"isolated_primary", "compound_primary"}
        ]
        prior_rows = [
            row for row in rows
            if row["cohort"] == cohort and row["record_type"] == "prior_primary_count"
        ]
        for column in ("n_episodes", "n_persons"):
            total = sum(Decimal(row[column]) for row in episode_rows if row[column])
            published = sum(Decimal(row[column]) for row in prior_rows if row[column])
            suppressed_residual = total - published
            if Decimal("1") <= suppressed_residual <= Decimal("10"):
                fail(
                    f"complementary disclosure in {episode_path.relative_to(REPO)}: "
                    f"{cohort} {column} suppressed residual={suppressed_residual}"
                )


def check_cross_file_sample_counts() -> None:
    """Reject within-sample count pairs whose subtraction exposes 1 through 10.

    Only explicitly named participant/sample count fields are considered. Positions,
    waves, proportions, tuning parameters and model estimates are deliberately excluded.
    Display-source denominators are linked back to sample_id through their provenance row.
    """
    observations: dict[str, list[tuple[int, str, str, int]]] = defaultdict(list)
    output_files = sorted((REPO / "source_data/analysis_outputs").rglob("*.csv"))
    cached: dict[Path, list[dict[str, str]]] = {}

    for path in output_files:
        rows = _records(path)
        cached[path] = rows
        for row_number, row in enumerate(rows, start=2):
            sample_id = str(row.get("sample_id", "")).strip()
            if not sample_id:
                continue
            for column in SAMPLE_PERSON_COUNT_COLUMNS & set(row):
                value = _integer_count(row.get(column, ""))
                if value is not None and value > 0:
                    observations[sample_id].append(
                        (value, path.relative_to(REPO).as_posix(), column, row_number)
                    )

    display_files = sorted((REPO / "source_data/main").glob("*.csv")) + sorted(
        (REPO / "source_data/extended_data").glob("*.csv")
    )
    for path in display_files:
        for row_number, row in enumerate(_records(path), start=2):
            source_text = str(row.get("source_file", "")).strip()
            source_row = _integer_count(row.get("source_row_0based", ""))
            if not source_text or source_row is None:
                continue
            source = REPO / source_text
            source_records = cached.get(source)
            if source_records is None:
                source_records = _records(source)
                cached[source] = source_records
            original = source_records[source_row]
            sample_id = str(original.get("sample_id", "")).strip()
            if not sample_id:
                continue
            for column in ("n", "n_complete"):
                value = _integer_count(row.get(column, ""))
                if value is not None and value > 0:
                    observations[sample_id].append(
                        (value, path.relative_to(REPO).as_posix(), column, row_number)
                    )

    for sample_id, records in observations.items():
        by_value: dict[int, list[tuple[str, str, int]]] = defaultdict(list)
        for value, path, column, row_number in records:
            by_value[value].append((path, column, row_number))
        ordered = sorted(by_value)
        for lower, upper in zip(ordered, ordered[1:]):
            difference = upper - lower
            if 1 <= difference <= 10:
                left = by_value[lower][0]
                right = by_value[upper][0]
                fail(
                    f"within-sample count subtraction exposes {difference}: sample_id={sample_id}; "
                    f"{left[0]}:{left[2]} {left[1]}={lower}; "
                    f"{right[0]}:{right[2]} {right[1]}={upper}"
                )


def check_public_language() -> None:
    forbidden = ("locally frozen", "reference-matched", "existing aggregate results", "candidate repository")
    paths = [REPO / "README.md", REPO / "CHANGELOG.md", *sorted((REPO / "docs").glob("*.md")), *sorted((REPO / "source_data/main").glob("*.csv")), *sorted((REPO / "source_data/extended_data").glob("*.csv"))]
    for path in paths:
        text = path.read_text(encoding="utf-8", errors="replace").lower()
        for phrase in forbidden:
            if phrase in text:
                fail(f"internal wording in {path.relative_to(REPO)}: {phrase}")


def check_hidden_metadata() -> None:
    ignored_parts = {".git", "__MACOSX", "__pycache__", "outputs"}
    for path in sorted(REPO.rglob("*")):
        if not path.is_file() or any(part in ignored_parts for part in path.relative_to(REPO).parts):
            continue
        rel = path.relative_to(REPO)
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            with zipfile.ZipFile(path) as archive:
                for member in archive.namelist():
                    if not member.lower().endswith((".xml", ".rels", ".txt")):
                        continue
                    text = archive.read(member).decode("utf-8", errors="replace")
                    if AI_TOOL_PATTERN.search(text):
                        fail(f"AI-tool name in Office metadata: {rel}:{member}")
                    if LOCAL_PATH_PATTERN.search(text):
                        fail(f"local path in Office metadata: {rel}:{member}")
                    if any(pattern.search(text) for pattern in SECRET_PATTERNS):
                        fail(f"possible credential in Office metadata: {rel}:{member}")
        elif suffix == ".pdf":
            reader = PdfReader(path)
            metadata = " ".join(str(value) for value in (reader.metadata or {}).values())
            extracted = " ".join(page.extract_text() or "" for page in reader.pages)
            combined = metadata + " " + extracted
            if AI_TOOL_PATTERN.search(combined):
                fail(f"AI-tool name in PDF metadata or text: {rel}")
            if LOCAL_PATH_PATTERN.search(combined):
                fail(f"local path in PDF metadata or text: {rel}")
            if any(pattern.search(combined) for pattern in SECRET_PATTERNS):
                fail(f"possible credential in PDF metadata or text: {rel}")


def check_workbooks() -> None:
    for path in sorted(REPO.rglob("*.xlsx")):
        book = load_workbook(path, read_only=False, data_only=False)
        if book._external_links:
            fail(f"external workbook link: {path.relative_to(REPO)}")
        for sheet in book.worksheets:
            rows = list(sheet.iter_rows())
            if rows:
                headers = {str(cell.value or "").strip().lower() for cell in rows[0]}
                found = sorted(headers & ID_HEADERS)
                if found:
                    fail(f"identifier fields in {path.relative_to(REPO)}:{sheet.title}: {found}")
            for row in rows:
                for cell in row:
                    if cell.data_type == "e" or (isinstance(cell.value, str) and cell.value.startswith(ERRORS)):
                        fail(f"spreadsheet error in {path.relative_to(REPO)}:{sheet.title}!{cell.coordinate}")
        book.close()


def check_role_correction() -> None:
    book = load_workbook(REPO / "source_data/workbooks/Source_Data_Figure_1.xlsx", read_only=True, data_only=False)
    if book["Panel_B"]["F2"].value != "primary cross-adversity analysis":
        fail("Figure 1 Panel_B!F2 role mismatch")
    if book["Panel_B"]["F3"].value != "independent conceptual replication":
        fail("Figure 1 Panel_B!F3 role mismatch")
    book.close()
    with (REPO / "source_data/main/Figure_1_source.csv").open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    role = {int(row["display_row"]): row["analysis_role"] for row in rows}
    if role[2] != "primary cross-adversity analysis" or role[7] != "independent conceptual replication":
        fail("Figure 1 CSV role metadata mismatch")



def check_registry() -> None:
    registry_path = REPO / "config/display_registry.json"
    registry = json.loads(registry_path.read_text(encoding="utf-8"))
    figures = registry.get("figures", {})
    tables = registry.get("tables", {})
    if len([key for key, value in figures.items() if value.get("main")]) != 4:
        fail("display registry must contain four main figures")
    if len([key for key in figures if key.startswith("Extended_Data_Figure_")]) != 8:
        fail("display registry must contain eight Extended Data figures")
    if len(tables) != 9:
        fail("display registry must contain nine table workbooks")
    for section, records in (("figures", figures), ("tables", tables)):
        for key, record in records.items():
            source = REPO / str(record.get("source", ""))
            if not source.is_file():
                fail(f"missing registry source for {section}.{key}: {source.relative_to(REPO)}")

def check_pdfs() -> None:
    expected = sorted((REPO / "expected_outputs").rglob("*.pdf"))
    if len(expected) != 12:
        fail(f"expected 12 reference figure PDFs, found {len(expected)}")
    for path in expected:
        reader = PdfReader(path)
        if len(reader.pages) != 1:
            fail(f"figure must be one page: {path.relative_to(REPO)}")
        metadata = {str(k): str(v) for k, v in (reader.metadata or {}).items()}
        joined = " ".join(metadata.values()).lower()
        if "anonymous" in joined or "reference-matched" in joined or "existing aggregate results" in joined or "no analysis" in joined:
            fail(f"internal PDF metadata: {path.relative_to(REPO)}")
        if metadata.get("/Author") != "Huiyun Yu and Bo Li":
            fail(f"missing figure author metadata: {path.relative_to(REPO)}")


def main() -> int:
    check_files()
    check_csv()
    check_display_provenance()
    check_small_counts()
    check_cross_file_sample_counts()
    check_public_language()
    check_hidden_metadata()
    check_workbooks()
    check_role_correction()
    check_registry()
    check_pdfs()
    print("PASS: public release structure, disclosure controls, provenance, workbooks and expected PDFs")
    return 0


if __name__ == "__main__":
    sys.exit(main())
