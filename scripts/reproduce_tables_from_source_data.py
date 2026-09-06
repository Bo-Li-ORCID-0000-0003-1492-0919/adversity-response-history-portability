#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import json
import shutil
from pathlib import Path

from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[1]
ERRORS = ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def validate(path: Path) -> None:
    book = load_workbook(path, read_only=False, data_only=False)
    if book._external_links:
        raise AssertionError(f"External links in {path}")
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "e" or (isinstance(cell.value, str) and cell.value.startswith(ERRORS)):
                    raise AssertionError(f"Spreadsheet error in {path}:{sheet.title}!{cell.coordinate}")
    book.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Validate public table workbooks and copy them to a separate output folder.")
    parser.add_argument("--output-dir", type=Path, default=REPO / "outputs/table_reproduction")
    args = parser.parse_args()
    sources = [
        *(REPO / "source_data/tables").glob("Table_*.xlsx"),
        *(REPO / "extended_data_tables").glob("Extended_Data_Table_*.xlsx"),
        *(REPO / "supplementary_tables").glob("Supplementary_Table_*.xlsx"),
    ]
    records = []
    for source in sorted(sources):
        validate(source)
        destination = args.output_dir / source.parent.name / source.name
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, destination)
        if sha(source) != sha(destination):
            raise AssertionError(f"Copy hash mismatch: {source.name}")
        records.append({"source": str(source.relative_to(REPO)), "output": str(destination), "sha256": sha(source)})
    args.output_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / "reproduction_manifest.json").write_text(json.dumps(records, indent=2) + "\n", encoding="utf-8")
    print(f"PASS: validated and copied {len(records)} aggregate table workbooks")


if __name__ == "__main__":
    main()
